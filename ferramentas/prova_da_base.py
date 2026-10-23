"""Gera o CONFIG.provaIdade do quiz a partir da planilha de avaliações das alunas.

Uso:
    python3 ferramentas/prova_da_base.py caminho/da/planilha.xlsx
    python3 ferramentas/prova_da_base.py caminho/da/planilha.xlsx --simulacao   # só para testar com base de exemplo

A planilha precisa da aba "Participantes" com as colunas:
    Faixa etária (18–29, 30–39, 40–49, 50–59, 60+), Usa caneta? (Sim/Não),
    Percepção da barriga (Sim / Um pouco / Não), Semanas com cardápio.

Regras:
- Conta só quem usa caneta (é o público do quiz).
- O número da frase é o % que respondeu "Sim" (não soma o "um pouco").
- Faixa com menos de MINIMO respostas fica de fora (o quiz mostra o dado de pesquisa nela).
- Se a planilha disser que é simulação, teste ou demonstração, não gera nada sem --simulacao.
"""
import json
import sys

import openpyxl

MINIMO = 30
FAIXAS = {'18–29': ('i18', '18 a 29 anos'), '30–39': ('i30', '30 a 39 anos'), '40–49': ('i40', '40 a 49 anos'),
          '50–59': ('i50', '50 a 59 anos'), '60+': ('i60', '60 anos ou mais')}
SINAIS_DE_TESTE = ('simulação', 'simulacao', 'teste', 'demonstração', 'demonstracao', 'exemplo')


def parece_teste(wb):
    for ws in wb.worksheets:
        if ws.title == 'Participantes':
            continue
        for linha in ws.iter_rows(values_only=True):
            for v in linha:
                if isinstance(v, str) and any(s in v.lower() for s in SINAIS_DE_TESTE):
                    return v
    return None


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    caminho, simulacao = sys.argv[1], '--simulacao' in sys.argv
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    aviso = parece_teste(wb)
    if aviso and not simulacao:
        sys.exit(f'PARADO: a planilha diz "{aviso}". Esses números não podem ir para o quiz no ar.\n'
                 'Para só testar a ferramenta, rode de novo com --simulacao.')

    linhas = list(wb['Participantes'].iter_rows(values_only=True))
    col = {nome: i for i, nome in enumerate(linhas[0])}
    for nome in ('Faixa etária', 'Usa caneta?', 'Percepção da barriga', 'Semanas com cardápio'):
        if nome not in col:
            sys.exit(f'Falta a coluna "{nome}" na aba Participantes.')

    contagem, semanas = {}, []
    for r in linhas[1:]:
        if r[col['Usa caneta?']] != 'Sim' or r[col['Faixa etária']] not in FAIXAS:
            continue
        c = contagem.setdefault(r[col['Faixa etária']], {'n': 0, 'sim': 0, 'pouco': 0})
        c['n'] += 1
        resp = r[col['Percepção da barriga']]
        c['sim'] += resp == 'Sim'
        c['pouco'] += resp == 'Um pouco'
        if isinstance(r[col['Semanas com cardápio']], (int, float)):
            semanas.append(r[col['Semanas com cardápio']])

    semanas.sort()
    mediana = round(semanas[len(semanas) // 2]) if semanas else None
    print(f'{"Faixa":<8}{"respostas":>10}{"% sim":>8}{"% sim+pouco":>13}')
    prova = {}
    for faixa, (chave, _) in FAIXAS.items():
        c = contagem.get(faixa)
        if not c:
            print(f'{faixa:<8}{0:>10}{"-":>8}{"-":>13}   sem respostas')
            prova[chave] = None
            continue
        pct, pct2 = round(100 * c['sim'] / c['n']), round(100 * (c['sim'] + c['pouco']) / c['n'])
        nota = '' if c['n'] >= MINIMO else f'   menos de {MINIMO} respostas: fica de fora'
        print(f'{faixa:<8}{c["n"]:>10}{pct:>7}%{pct2:>12}%{nota}')
        prova[chave] = None if c['n'] < MINIMO else {
            'pct': pct,
            'texto': 'que usam a caneta disseram que a barriga diminuiu' + (f' em cerca de {mediana} semanas de cardápio' if mediana else ' com o cardápio'),
        }

    print('\nCole no CONFIG do quiz/index.html (troque a linha provaIdade):\n')
    print('  provaIdade: ' + json.dumps(prova, ensure_ascii=False).replace('"pct"', 'pct').replace('"texto"', 'texto') + ',')
    if simulacao:
        print('\nATENÇÃO: base de simulação. Use só para ver como fica; não publique.')


if __name__ == '__main__':
    main()
